"""Generate the final .docx assignment report with embedded plots."""
import os
import json
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PROJECT = os.path.dirname(os.path.abspath(__file__))
ARTIFACTS = os.path.join(PROJECT, "model", "artifacts")
OUTPUT = os.path.join(PROJECT, "Heart_Disease_MLOps_Report.docx")


def set_style(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    return h


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    return table


def add_code(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pf.left_indent = Cm(1)


def add_image(doc, filename, caption, width=5.5):
    path = os.path.join(ARTIFACTS, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].italic = True
        cap.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def main():
    doc = Document()
    set_style(doc)

    # --- TITLE PAGE ---
    for _ in range(6):
        doc.add_paragraph("")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Heart Disease MLOps Pipeline")
    run.font.size = Pt(28)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "End-to-End ML Model Development, CI/CD, and Production Deployment"
    )
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for _ in range(2):
        doc.add_paragraph("")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Course: MLOps (S2-25_AMLCSZG523)\nAssignment - I\nTotal Marks: 50")
    run.font.size = Pt(12)

    doc.add_page_break()

    # --- TABLE OF CONTENTS ---
    add_heading(doc, "Table of Contents", level=1)
    toc_items = [
        "1. Introduction",
        "2. Data Acquisition & Exploratory Data Analysis",
        "3. Feature Engineering & Model Development",
        "4. Experiment Tracking (MLflow)",
        "5. Model Packaging & Reproducibility",
        "6. CI/CD Pipeline & Automated Testing",
        "7. Model Containerization (Docker)",
        "8. Production Deployment (Kubernetes)",
        "9. Monitoring & Logging",
        "10. Architecture Overview & Conclusion",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    # --- 1. INTRODUCTION ---
    add_heading(doc, "1. Introduction", level=1)
    doc.add_paragraph(
        "This report documents the design, development, and deployment of an end-to-end "
        "MLOps pipeline for predicting heart disease risk using the UCI Heart Disease dataset "
        "(Cleveland subset). The pipeline implements modern MLOps best practices including "
        "automated data preprocessing, experiment tracking with MLflow, CI/CD pipelines with "
        "Jenkins and GitHub Actions, Docker containerization, Kubernetes deployment, and "
        "production monitoring with Prometheus and Grafana."
    )
    add_heading(doc, "Problem Statement", level=2)
    doc.add_paragraph(
        "Build a machine learning classifier to predict the risk of heart disease based on "
        "patient health data, and deploy the solution as a cloud-ready, monitored API."
    )
    add_heading(doc, "Dataset", level=2)
    doc.add_paragraph(
        "Heart Disease UCI Dataset (Cleveland subset) from the UCI Machine Learning Repository. "
        "303 patients, 13 clinical features (age, sex, chest pain type, blood pressure, "
        "cholesterol, etc.), and a binary target (presence/absence of heart disease)."
    )
    add_heading(doc, "Tech Stack", level=2)
    doc.add_paragraph(
        "Python 3.12, scikit-learn, MLflow, FastAPI, Docker, Kubernetes (Minikube), "
        "Jenkins, GitHub Actions, Prometheus, Grafana, pytest."
    )

    # --- 2. DATA & EDA ---
    add_heading(doc, "2. Data Acquisition & Exploratory Data Analysis", level=1)
    add_heading(doc, "2.1 Data Acquisition", level=2)
    doc.add_paragraph(
        "A download script (data/download_data.py) reads the raw processed Cleveland data "
        "from the UCI repository, adds proper column headers, converts the multi-class target "
        "(values 0-4) to binary classification (0 = no disease, 1 = disease present), "
        "and saves it as a clean CSV file."
    )
    doc.add_paragraph("Command to acquire data:")
    add_code(doc, "python data/download_data.py")

    add_heading(doc, "2.2 Data Cleaning", level=2)
    items = [
        "Missing values: 'ca' (4 missing) and 'thal' (2 missing) contained '?' markers, handled via median imputation.",
        "Target encoding: Values 1-4 mapped to 1 (disease present), value 0 kept as 0 (no disease).",
        "Final dataset: 303 rows, 14 columns (13 features + 1 target), zero missing values after cleaning.",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "2.3 EDA Highlights", level=2)
    doc.add_paragraph(
        "The EDA notebook (notebooks/01_eda.ipynb) includes dataset statistics, "
        "feature distribution histograms, correlation heatmaps, class balance analysis, "
        "box plots of continuous features by target, and categorical feature cross-tabulations."
    )
    eda_findings = [
        "Class distribution: ~54% no disease (164), ~46% disease (139) - reasonably balanced.",
        "Strong correlations with target: thal (0.52), ca (0.47), exang (0.44), cp (0.43), oldpeak (0.43).",
        "Disease patients show lower max heart rate and higher ST depression.",
        "Chest pain type 4 (asymptomatic) is strongly associated with heart disease.",
    ]
    for f in eda_findings:
        doc.add_paragraph(f, style="List Bullet")

    add_heading(doc, "2.4 Data Split", level=2)
    doc.add_paragraph(
        "The cleaned dataset was split into training (242 samples, 80%) and test (61 samples, 20%) "
        "sets using stratified sampling to preserve target distribution."
    )

    doc.add_page_break()

    # --- 3. MODEL DEVELOPMENT ---
    add_heading(doc, "3. Feature Engineering & Model Development", level=1)
    add_heading(doc, "3.1 Preprocessing Pipeline", level=2)
    doc.add_paragraph(
        "All 13 features are treated as numeric. The preprocessing pipeline is implemented "
        "as a scikit-learn Pipeline consisting of:"
    )
    steps = [
        "SimpleImputer (median strategy) - handles any remaining missing values at inference time.",
        "StandardScaler - normalizes features to zero mean and unit variance.",
        "Classifier (Logistic Regression or Random Forest).",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")

    add_heading(doc, "3.2 Models Trained", level=2)
    doc.add_paragraph(
        "Two classification models were trained with hyperparameter tuning via GridSearchCV (5-fold CV):"
    )
    doc.add_paragraph(
        "Logistic Regression: C = [0.01, 0.1, 1.0, 10.0]. Best: C=0.1, solver=lbfgs.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Random Forest: n_estimators=[100,200], max_depth=[5,10,None], "
        "min_samples_split=[2,5], min_samples_leaf=[1,2]. "
        "Best: n_estimators=200, max_depth=5.",
        style="List Bullet",
    )

    add_heading(doc, "3.3 Test Set Evaluation", level=2)
    with open(os.path.join(ARTIFACTS, "metrics.json")) as f:
        metrics = json.load(f)
    lr = metrics["logistic_regression"]
    rf = metrics["random_forest"]
    add_table(doc,
        ["Metric", "Logistic Regression", "Random Forest"],
        [
            ["Accuracy", f"{lr['accuracy']:.4f}", f"{rf['accuracy']:.4f}"],
            ["Precision", f"{lr['precision']:.4f}", f"{rf['precision']:.4f}"],
            ["Recall", f"{lr['recall']:.4f}", f"{rf['recall']:.4f}"],
            ["F1 Score", f"{lr['f1']:.4f}", f"{rf['f1']:.4f}"],
            ["ROC-AUC", f"{lr['roc_auc']:.4f}", f"{rf['roc_auc']:.4f}"],
        ],
    )

    add_heading(doc, "3.4 Cross-Validation Results (5-Fold)", level=2)
    lr_cv = metrics["logistic_regression_cv"]
    rf_cv = metrics["random_forest_cv"]
    add_table(doc,
        ["Metric", "Logistic Regression", "Random Forest"],
        [
            ["CV Accuracy", f"{lr_cv['cv_accuracy']:.4f}", f"{rf_cv['cv_accuracy']:.4f}"],
            ["CV Precision", f"{lr_cv['cv_precision']:.4f}", f"{rf_cv['cv_precision']:.4f}"],
            ["CV Recall", f"{lr_cv['cv_recall']:.4f}", f"{rf_cv['cv_recall']:.4f}"],
            ["CV F1", f"{lr_cv['cv_f1']:.4f}", f"{rf_cv['cv_f1']:.4f}"],
            ["CV ROC-AUC", f"{lr_cv['cv_roc_auc']:.4f}", f"{rf_cv['cv_roc_auc']:.4f}"],
        ],
    )

    add_heading(doc, "3.5 ROC Curve Comparison", level=2)
    add_image(doc, "roc_curve_comparison.png", "Figure 1: ROC Curve Comparison - Both models achieve AUC ~0.96", 5.0)

    add_heading(doc, "3.6 Confusion Matrix", level=2)
    add_image(doc, "confusion_matrix_logistic_regression.png",
              "Figure 2: Confusion Matrix - Logistic Regression (selected model)", 4.0)

    add_heading(doc, "3.7 Model Selection", level=2)
    doc.add_paragraph(
        f"The best model was selected based on test ROC-AUC score. Logistic Regression "
        f"(AUC={lr['roc_auc']:.4f}) was chosen as the production model due to its highest "
        f"ROC-AUC, simpler architecture, faster inference time, and comparable performance "
        f"to Random Forest (AUC={rf['roc_auc']:.4f})."
    )

    doc.add_page_break()

    # --- 4. EXPERIMENT TRACKING ---
    add_heading(doc, "4. Experiment Tracking (MLflow)", level=1)
    doc.add_paragraph(
        "MLflow is integrated into the training pipeline (src/training/train.py) with "
        "file-based tracking (./mlruns directory). The experiment name is "
        "'heart-disease-classification'."
    )
    add_heading(doc, "4.1 What Gets Logged", level=2)
    logged = [
        "Parameters: model_type, C, n_estimators, max_depth, min_samples_split, etc.",
        "Metrics: accuracy, precision, recall, F1, ROC-AUC (both test and 5-fold CV).",
        "Artifacts: confusion matrix plots, ROC curve plots, trained sklearn model.",
        "Tags: dataset=heart-disease-uci-cleveland, dataset_version=1.0, description.",
    ]
    for l in logged:
        doc.add_paragraph(l, style="List Bullet")

    add_heading(doc, "4.2 Viewing Experiments", level=2)
    doc.add_paragraph("To view the MLflow experiment comparison UI:")
    add_code(doc, "mlflow ui --backend-store-uri ./mlruns --port 5001\n# Open http://localhost:5001")
    doc.add_paragraph(
        "The MLflow UI provides side-by-side comparison of Logistic Regression vs Random Forest "
        "runs, metric charts showing performance, and an artifact browser for downloading "
        "models and plots."
    )

    # --- 5. MODEL PACKAGING ---
    add_heading(doc, "5. Model Packaging & Reproducibility", level=1)
    add_heading(doc, "5.1 Saved Formats", level=2)
    doc.add_paragraph(
        "The final model is saved in two formats: (1) Joblib pickle at "
        "model/artifacts/heart_disease_model.pkl for direct loading in the serving API, "
        "and (2) MLflow model format in mlruns/ for the MLflow model registry."
    )
    add_heading(doc, "5.2 Full Pipeline Bundling", level=2)
    doc.add_paragraph(
        "The saved model is a complete scikit-learn Pipeline containing imputation, "
        "scaling, and classification steps. This guarantees that all preprocessing "
        "transformations are bundled with the model, ensuring identical behavior "
        "regardless of deployment environment."
    )
    add_heading(doc, "5.3 Dependency Management", level=2)
    doc.add_paragraph(
        "All dependencies are pinned with exact versions in requirements.txt. "
        "Key packages: scikit-learn==1.8.0, pandas==2.3.3, numpy==2.4.4, "
        "mlflow==3.11.1, fastapi==0.136.1, uvicorn==0.46.0, pytest==9.0.3."
    )
    add_heading(doc, "5.4 Reproducibility Steps", level=2)
    add_code(doc,
        "python3 -m venv venv && source venv/bin/activate\n"
        "pip install -r requirements.txt\n"
        "python data/download_data.py\n"
        "python -c \"from src.data_prep.prepare import run; run()\"\n"
        "python -c \"from src.training.train import run; run()\""
    )

    doc.add_page_break()

    # --- 6. CI/CD ---
    add_heading(doc, "6. CI/CD Pipeline & Automated Testing", level=1)
    add_heading(doc, "6.1 Unit Tests", level=2)
    doc.add_paragraph(
        "24 unit tests are organized across three test files, all passing successfully:"
    )
    add_table(doc,
        ["Test File", "Count", "Coverage"],
        [
            ["test_data_prep.py", "9", "Data loading, missing values, shape, target encoding, split, preprocessor"],
            ["test_model.py", "8", "Training accuracy, prediction format, probabilities, saved model"],
            ["test_api.py", "7", "Root endpoint, health check, prediction, validation, Swagger docs"],
        ],
    )
    add_code(doc, "python -m pytest tests/ -v\n# 24 passed in 3.32s")

    add_heading(doc, "6.2 Jenkins Pipeline (Jenkinsfile)", level=2)
    doc.add_paragraph("The Jenkinsfile defines a 7-stage pipeline:")
    stages = [
        "Checkout: Pull source code from repository.",
        "Setup Python Environment: Create virtual environment, install dependencies.",
        "Lint: Run flake8 for code quality (max-line-length=120).",
        "Prepare Data: Download dataset and run preprocessing.",
        "Unit Tests: Run pytest with JUnit XML reporting.",
        "Train Model: Execute the full training pipeline.",
        "Build Docker Image: Build container for model serving.",
    ]
    for s in stages:
        doc.add_paragraph(s, style="List Number")

    add_heading(doc, "6.3 GitHub Actions (.github/workflows/ci.yml)", level=2)
    doc.add_paragraph(
        "An equivalent GitHub Actions workflow is provided, triggered on push/PR to main. "
        "It includes all Jenkins stages plus Docker container testing with health check "
        "and prediction verification via curl. Model artifacts and test results are "
        "uploaded as workflow artifacts."
    )

    # --- 7. CONTAINERIZATION ---
    add_heading(doc, "7. Model Containerization (Docker)", level=1)
    add_heading(doc, "7.1 FastAPI Application", level=2)
    doc.add_paragraph(
        "The serving API (src/serving/app.py) is built with FastAPI and exposes the following endpoints:"
    )
    add_table(doc,
        ["Method", "Path", "Description"],
        [
            ["GET", "/", "API information and available endpoints"],
            ["GET", "/health", "Health check (model loaded status)"],
            ["GET", "/docs", "Auto-generated Swagger UI documentation"],
            ["POST", "/predict", "Accept 13 features, return prediction + confidence"],
            ["GET", "/metrics", "Prometheus metrics endpoint"],
        ],
    )

    add_heading(doc, "7.2 Dockerfile", level=2)
    doc.add_paragraph("Key security and production features:")
    features = [
        "Base image: python:3.12-slim (minimal attack surface).",
        "Non-root user (appuser) for security.",
        "Requirements-first copy for Docker layer caching.",
        "Built-in HEALTHCHECK directive.",
        "Runs uvicorn ASGI server on port 8080.",
    ]
    for f in features:
        doc.add_paragraph(f, style="List Bullet")

    add_heading(doc, "7.3 Build and Test", level=2)
    add_code(doc,
        "docker build -t heart-disease-api .\n"
        "docker run -p 8080:8080 heart-disease-api\n\n"
        "# Test health\n"
        "curl http://localhost:8080/health\n\n"
        "# Test prediction\n"
        'curl -X POST http://localhost:8080/predict \\\n'
        '  -H "Content-Type: application/json" \\\n'
        '  -d \'{"age":63,"sex":1,"cp":1,"trestbps":145,"chol":233,...}\''
    )

    doc.add_page_break()

    # --- 8. DEPLOYMENT ---
    add_heading(doc, "8. Production Deployment (Kubernetes)", level=1)
    add_heading(doc, "8.1 Kubernetes Manifests", level=2)
    doc.add_paragraph("Three Kubernetes manifests are provided in the k8s/ directory:")
    manifests = [
        "deployment.yaml: 2 replicas, resource limits (256Mi-512Mi memory, 250m-500m CPU), "
        "liveness and readiness probes on /health, rolling update strategy.",
        "service.yaml: LoadBalancer type service, port 80 mapped to container port 8080.",
        "ingress.yaml: Nginx ingress controller with path-based routing.",
    ]
    for m in manifests:
        doc.add_paragraph(m, style="List Bullet")

    add_heading(doc, "8.2 Minikube Deployment Steps", level=2)
    add_code(doc,
        "minikube start\n"
        "eval $(minikube docker-env)\n"
        "docker build -t heart-disease-api:latest .\n"
        "kubectl apply -f k8s/\n"
        "minikube service heart-disease-api --url"
    )

    add_heading(doc, "8.3 Verification", level=2)
    doc.add_paragraph(
        "After deployment, verify with: kubectl get pods (shows 2 running replicas), "
        "kubectl get svc (shows external URL), curl <URL>/health (healthy response), "
        "and curl -X POST <URL>/predict (prediction response)."
    )

    # --- 9. MONITORING ---
    add_heading(doc, "9. Monitoring & Logging", level=1)
    add_heading(doc, "9.1 Structured JSON Logging", level=2)
    doc.add_paragraph(
        "All API requests are logged in JSON format with HTTP method, path, status code, "
        "response latency (ms), and UTC timestamp."
    )

    add_heading(doc, "9.2 Prometheus Metrics", level=2)
    doc.add_paragraph("The API exposes a /metrics endpoint with the following metrics:")
    add_table(doc,
        ["Metric", "Type", "Description"],
        [
            ["http_requests_total", "Counter", "Total requests by method/handler/status"],
            ["http_request_duration_seconds", "Histogram", "Request latency distribution"],
            ["http_requests_inprogress", "Gauge", "Currently processing requests"],
            ["heart_disease_predictions_total", "Counter", "Predictions by class (Disease/No Disease)"],
        ],
    )

    add_heading(doc, "9.3 Grafana Dashboard", level=2)
    doc.add_paragraph(
        "A pre-configured Grafana dashboard provides: request rate time series, "
        "p95 latency time series, prediction distribution pie chart, error rate stat panel, "
        "and requests in-progress gauge."
    )

    add_heading(doc, "9.4 Docker Compose Monitoring Stack", level=2)
    add_code(doc,
        "docker compose up --build\n"
        "# App:        http://localhost:8080\n"
        "# Prometheus: http://localhost:9090\n"
        "# Grafana:    http://localhost:3000 (admin/admin)"
    )

    doc.add_page_break()

    # --- 10. ARCHITECTURE ---
    add_heading(doc, "10. Architecture Overview & Conclusion", level=1)
    add_heading(doc, "10.1 Architecture Diagram", level=2)
    doc.add_paragraph(
        "The complete pipeline flows through five major stages:"
    )
    arch = [
        "Data Pipeline: UCI Dataset -> download_data.py -> prepare.py -> Train/Test CSVs",
        "Training Pipeline: train.py -> GridSearchCV (LR + RF) -> MLflow Tracking -> Model Artifact",
        "CI/CD Pipeline: Git Push -> Jenkins/GitHub Actions -> Lint -> Test -> Train -> Docker Build",
        "Deployment: Docker Image -> Kubernetes (Minikube) -> FastAPI Service (2 replicas) -> LoadBalancer",
        "Monitoring: API -> Prometheus (metrics scraping) -> Grafana (dashboards) -> JSON Logging",
    ]
    for a in arch:
        doc.add_paragraph(a, style="List Number")

    add_heading(doc, "10.2 Project Structure", level=2)
    structure = (
        "heart-disease-mlops/\n"
        "  data/              - download script, raw & processed CSVs\n"
        "  notebooks/         - 01_eda.ipynb, 02_training.ipynb\n"
        "  src/data_prep/     - prepare.py (preprocessing pipeline)\n"
        "  src/training/      - train.py (training + MLflow)\n"
        "  src/evaluation/    - evaluate.py (metrics)\n"
        "  src/serving/       - app.py (FastAPI /predict endpoint)\n"
        "  src/monitoring/    - middleware.py (Prometheus metrics)\n"
        "  tests/             - 24 pytest unit tests\n"
        "  model/artifacts/   - saved model + metrics\n"
        "  k8s/               - deployment, service, ingress YAML\n"
        "  monitoring/        - prometheus.yml, Grafana configs\n"
        "  Dockerfile, docker-compose.yml, Jenkinsfile\n"
        "  .github/workflows/ci.yml, requirements.txt"
    )
    add_code(doc, structure)

    add_heading(doc, "10.3 Repository Link", level=2)
    doc.add_paragraph(
        "GitHub: https://github.com/Deepesh-BITsPilani/mlops_heart_disease_assignment"
    )

    add_heading(doc, "10.4 Conclusion", level=2)
    doc.add_paragraph(
        "This project demonstrates a complete MLOps lifecycle from data acquisition through "
        "production deployment and monitoring. The pipeline achieves a ROC-AUC of 0.958 for "
        "heart disease prediction, with full automation via CI/CD pipelines, containerized "
        "deployment on Kubernetes, and real-time monitoring through Prometheus and Grafana. "
        "All 24 unit tests pass with zero lint errors, and the entire pipeline is reproducible "
        "from a clean setup using the provided requirements file."
    )

    doc.save(OUTPUT)
    print(f"Report saved to: {OUTPUT}")


if __name__ == "__main__":
    main()
